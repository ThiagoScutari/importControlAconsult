"""Item 9 [F1-09] — Saída D: append na Relação das Importações (.xlsx).

Usa o arquivo REAL da empresa (assets) como fixture: a ferramenta insere UMA
linha sob o bloco do ano, preserva as fórmulas das linhas existentes e as demais
abas, e o arquivo continua abrindo (recarregável pelo openpyxl = XML íntegro).
"""
import io
import os

import openpyxl
import pytest
from fastapi.testclient import TestClient

from backend.main import app
from backend.models import MiddlewareInput, ProcessoExtraido
from backend.outputs import append_relacao

client = TestClient(app)

RELACAO_XLSX = "CONTROLE IMPORTAÇÕES - FORNECEDORES ESTRANGEIROS.xlsx"


@pytest.fixture
def relacao_bytes(assets_dir):
    caminho = os.path.join(assets_dir, RELACAO_XLSX)
    with open(caminho, "rb") as fh:
        return fh.read()


def _proc_2025():
    return ProcessoExtraido(
        tipo_importacao="PRÓPRIA",
        data_nf="15/09/2025",
        di_duimp="25/9998887-0",
        numero_nf="000.005.123",
        valor_nf=350000.50,
        processo="1159",
        despachante="SYNDEX",
        fornecedor_estrangeiro="ACME INDUSTRIAL CO LTD",
        invoice_usd=60000.0,
        tx_di=5.28,
    )


def test_append_insere_sob_bloco_do_ano_sem_corromper(relacao_bytes):
    wb0 = openpyxl.load_workbook(io.BytesIO(relacao_bytes))
    abas0 = list(wb0.sheetnames)
    ws0 = wb0["Relação importações"]
    h3_original = ws0["H3"].value  # fórmula da 1ª linha de dados

    out = append_relacao(relacao_bytes, _proc_2025(), MiddlewareInput(conta_processo_numero="1648"))

    wb = openpyxl.load_workbook(io.BytesIO(out))  # recarrega -> XML íntegro (abre no Excel)
    assert list(wb.sheetnames) == abas0            # todas as abas preservadas
    ws = wb["Relação importações"]

    # linha existente NÃO foi tocada (insert acima dela nunca acontece)
    assert ws["H3"].value == h3_original

    # nova linha logo abaixo do bloco 2025 (R4)
    assert ws["A4"].value == "PRÓPRIA"
    assert ws["C4"].value == "25/9998887-0"
    assert ws["D4"].value == 5123               # [R3] Nº NF inteiro
    assert ws["M4"].value == "1159"


def test_append_rebaseia_formulas_calculadas(relacao_bytes):
    out = append_relacao(relacao_bytes, _proc_2025(), None)
    ws = openpyxl.load_workbook(io.BytesIO(out))["Relação importações"]
    assert ws["H4"].value == '=CONCATENATE(I4," ",M4," ",J4," ",C4," ",K4," ",D4)'
    assert ws["R4"].value == "=P4*Q4"           # col 3 = invoice × tx DI
    assert ws["V4"].value == "=T4*U4"           # col 6
    assert ws["W4"].value == "=T4*Q4"           # col 7
    assert ws["X4"].value == "=V4-W4"           # col 8 = variação


def test_append_cria_novo_bloco_para_ano_inexistente(relacao_bytes):
    proc = _proc_2025()
    proc.data_nf = "20/03/2026"
    proc.di_duimp = "26BR0000258971-1"   # número BR -> rótulo DUIMP
    proc.processo = "1453"
    out = append_relacao(relacao_bytes, proc, None)
    ws = openpyxl.load_workbook(io.BytesIO(out))["Relação importações"]

    marcador = next(
        r for r in range(1, ws.max_row + 1)
        if (ws.cell(r, 1).value in (None, "")) and ws.cell(r, 2).value == 2026
    )
    dados = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 13).value == "1453")
    assert dados > marcador
    assert ws.cell(dados, 10).value == "DUIMP"   # col J: rótulo derivado do número


def _linha_sintetica(ws, r, tipo, di, nf, proc, invoice, tx, usd_pg, tx_cbo):
    ws[f"A{r}"], ws[f"C{r}"], ws[f"D{r}"], ws[f"M{r}"] = tipo, di, nf, proc
    ws[f"I{r}"], ws[f"J{r}"], ws[f"K{r}"] = "PROCESSO", "DI", "NF"
    ws[f"P{r}"], ws[f"Q{r}"], ws[f"T{r}"], ws[f"U{r}"] = invoice, tx, usd_pg, tx_cbo
    ws[f"H{r}"] = f'=CONCATENATE(I{r}," ",M{r}," ",J{r}," ",C{r}," ",K{r}," ",D{r})'
    ws[f"R{r}"], ws[f"V{r}"] = f"=P{r}*Q{r}", f"=T{r}*U{r}"
    ws[f"W{r}"], ws[f"X{r}"] = f"=T{r}*Q{r}", f"=V{r}-W{r}"


def _wb_2025_acima_de_2026() -> bytes:
    """Planilha sintética: bloco 2025 (R1..R3) ACIMA do bloco 2026 (R4..R5)."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Relação importações"
    ws["B1"] = 2025                                   # marcador 2025
    ws["A2"] = "TIPO IMPORTAÇÃO"                      # cabeçalho
    _linha_sintetica(ws, 3, "PRÓPRIA", "25/1111111-0", 100, "700", 10000, 5.0, 10000, 4.9)
    ws["B4"] = 2026                                   # marcador 2026 (abaixo do 2025)
    _linha_sintetica(ws, 5, "CONTA E ORDEM", "26/2222222-0", 200, "800", 20000, 5.2, 20000, 5.1)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_insercao_em_2025_desloca_2026_sem_corromper():
    """Insere linha no bloco 2025 (que está ACIMA do 2026): o bloco 2026 desce
    inteiro, sem sobrescrita e com as fórmulas reajustadas para a nova posição."""
    original = _wb_2025_acima_de_2026()
    snap_2026 = {}
    ws0 = openpyxl.load_workbook(io.BytesIO(original))["Relação importações"]
    for L in ("A", "C", "D", "M", "P", "Q", "T", "U"):
        snap_2026[L] = ws0[f"{L}5"].value

    out = append_relacao(original, _proc_2025(), None)
    ws = openpyxl.load_workbook(io.BytesIO(out))["Relação importações"]

    # bloco 2026 desceu (marcador e dados uma linha abaixo)
    d2026 = next(r for r in range(1, ws.max_row + 1) if ws.cell(r, 13).value == "800")
    m2026 = next(r for r in range(1, ws.max_row + 1)
                 if ws.cell(r, 1).value in (None, "") and ws.cell(r, 2).value == 2026)
    assert m2026 < d2026

    # valores de entrada do 2026 NÃO foram sobrescritos
    for L, v in snap_2026.items():
        assert ws[f"{L}{d2026}"].value == v, f"2026 col {L} sobrescrita"

    # fórmulas do 2026 rebaseadas para a nova linha (não apontam mais p/ a de cima)
    assert ws[f"R{d2026}"].value == f"=P{d2026}*Q{d2026}"
    assert ws[f"X{d2026}"].value == f"=V{d2026}-W{d2026}"
    assert ws[f"H{d2026}"].value == (
        f'=CONCATENATE(I{d2026}," ",M{d2026}," ",J{d2026}," ",C{d2026}," ",K{d2026}," ",D{d2026})'
    )

    # R3 (linha 2025 original) intacta e a nova linha 2025 entrou em R4
    assert ws["R3"].value == "=P3*Q3"
    assert ws["M4"].value == "1159"
    assert ws["D4"].value == 5123               # Nº NF de _proc_2025 (000.005.123)


def test_coluna_G_texto_igual_ao_resultado_da_formula_H(relacao_bytes):
    """A coluna G (texto) reproduz exatamente o que a fórmula da coluna H concatena."""
    out = append_relacao(relacao_bytes, _proc_2025(), None)
    ws = openpyxl.load_workbook(io.BytesIO(out))["Relação importações"]
    i, m, j, c, k, d = (ws[f"{L}4"].value for L in ("I", "M", "J", "C", "K", "D"))
    resultado_h = f"{i} {m} {j} {c} {k} {d}"     # == CONCATENATE(I," ",M," ",J," ",C," ",K," ",D)
    assert ws["G4"].value == resultado_h


def test_rota_relacao_devolve_xlsx_atualizado_com_nome_diferente(relacao_bytes):
    proc = _proc_2025()
    payload = {"processo": proc.model_dump(), "middleware": {"conta_processo_numero": "1648"}}
    resp = client.post(
        "/relacao",
        files={"arquivo": (RELACAO_XLSX, relacao_bytes,
                           "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"payload": __import__("json").dumps(payload)},
    )
    assert resp.status_code == 200
    assert "spreadsheetml" in resp.headers["content-type"]
    # o arquivo baixado tem nome DIFERENTE do original (não sobrescreve o upload)
    disp = resp.headers["content-disposition"]
    assert RELACAO_XLSX not in disp
    assert "relacao_atualizada" in disp
    ws = openpyxl.load_workbook(io.BytesIO(resp.content))["Relação importações"]
    assert ws["M4"].value == "1159"              # a linha do processo entrou
